"""
DOCX Loader
===========
Load Microsoft Word documents.
"""

from pathlib import Path
from typing import List

from .base import BaseLoader, Document


class DocxLoader(BaseLoader):
    """
    Loader for Microsoft Word documents.
    
    Uses python-docx for text extraction.
    Supports: .docx
    """
    
    supported_extensions: List[str] = [".docx"]
    
    def __init__(self, include_tables: bool = True):
        """
        Initialize the DOCX loader.
        
        Args:
            include_tables: Whether to include table content
        """
        self.include_tables = include_tables
    
    def load(self, file_path: str) -> Document:
        """
        Load a DOCX file.
        
        Args:
            file_path: Path to the DOCX file
            
        Returns:
            Document with extracted text and metadata
        """
        try:
            import docx
        except ImportError:
            raise ImportError("python-docx is required for DOCX loading. Install with: pip install python-docx")
        
        path = self._validate_file(file_path)
        
        # Read document
        try:
            doc = docx.Document(str(path))
        except Exception as e:
            raise ValueError(f"Error reading DOCX {file_path}: {e}")
        
        # Extract paragraphs
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        
        # Extract tables if requested
        tables_text = []
        if self.include_tables:
            for table in doc.tables:
                table_content = self._extract_table(table)
                if table_content:
                    tables_text.append(table_content)
        
        # Combine content
        content_parts = paragraphs
        if tables_text:
            content_parts.extend(tables_text)
        
        content = "\n\n".join(content_parts)
        
        # Create metadata
        metadata = self._create_metadata(path)
        metadata["format"] = "docx"
        metadata["paragraph_count"] = len(paragraphs)
        metadata["table_count"] = len(doc.tables)
        
        # Extract core properties if available
        core = doc.core_properties
        if core.title:
            metadata["title"] = core.title
        if core.author:
            metadata["author"] = core.author
        if core.subject:
            metadata["subject"] = core.subject
        if core.created:
            metadata["doc_created"] = core.created.isoformat()
        
        return Document(content=content, metadata=metadata)
    
    def _extract_table(self, table) -> str:
        """
        Extract text from a table.
        
        Args:
            table: docx Table object
            
        Returns:
            Formatted table text
        """
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        
        if rows:
            return "\n".join(rows)
        return ""
